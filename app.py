"""ATS Resume Checker - Streamlit + Google Gemini Flash.

Upload a resume (PDF / DOCX / TXT), optionally paste a job description,
and get an ATS score with concrete improvement suggestions.
"""

import io
import json
import os
import re

import streamlit as st
from docx import Document
from google import genai
from google.genai import types
from pypdf import PdfReader

DEFAULT_MODEL = "gemini-flash-latest"  # alias that always points to the newest Flash
MAX_FILE_MB = 5
MAX_RESUME_CHARS = 20_000
MIN_RESUME_CHARS = 150

SYSTEM_PROMPT = """You are an expert ATS (Applicant Tracking System) analyst and \
professional resume reviewer. Evaluate the resume text strictly and honestly. \
Do not inflate scores. Base every finding only on the text provided. \
If a job description is supplied, judge keyword and skill match against it; \
otherwise judge against general best practices for the candidate's apparent field. \
Respond ONLY with JSON matching the requested schema."""

RESPONSE_SCHEMA = {
    "type": "object",
    "properties": {
        "overall_score": {"type": "integer", "description": "0-100"},
        "summary": {"type": "string"},
        "category_scores": {
            "type": "object",
            "properties": {
                "formatting": {"type": "integer"},
                "keywords": {"type": "integer"},
                "experience_impact": {"type": "integer"},
                "skills": {"type": "integer"},
                "education": {"type": "integer"},
                "readability": {"type": "integer"},
            },
            "required": [
                "formatting",
                "keywords",
                "experience_impact",
                "skills",
                "education",
                "readability",
            ],
        },
        "strengths": {"type": "array", "items": {"type": "string"}},
        "missing_keywords": {"type": "array", "items": {"type": "string"}},
        "improvements": {
            "type": "array",
            "items": {
                "type": "object",
                "properties": {
                    "priority": {"type": "string", "description": "High, Medium or Low"},
                    "issue": {"type": "string"},
                    "suggestion": {"type": "string"},
                },
                "required": ["priority", "issue", "suggestion"],
            },
        },
        "rewrite_examples": {
            "type": "array",
            "items": {
                "type": "object",
                "properties": {
                    "original": {"type": "string"},
                    "improved": {"type": "string"},
                },
                "required": ["original", "improved"],
            },
        },
    },
    "required": [
        "overall_score",
        "summary",
        "category_scores",
        "strengths",
        "missing_keywords",
        "improvements",
        "rewrite_examples",
    ],
}

CATEGORY_LABELS = {
    "formatting": "Formatting",
    "keywords": "Keywords",
    "experience_impact": "Experience & Impact",
    "skills": "Skills",
    "education": "Education",
    "readability": "Readability",
}


# --------------------------------------------------------------------------
# File parsing
# --------------------------------------------------------------------------
def extract_text(filename: str, data: bytes) -> str:
    """Extract plain text from a PDF, DOCX or TXT file's bytes."""
    name = filename.lower()
    if name.endswith(".pdf"):
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception:
                raise ValueError("This PDF is password-protected.")
        pages = [(page.extract_text() or "") for page in reader.pages]
        text = "\n".join(pages)
    elif name.endswith(".docx"):
        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    parts.append(cell.text)
        text = "\n".join(parts)
    elif name.endswith(".txt"):
        text = data.decode("utf-8", errors="ignore")
    else:
        raise ValueError("Unsupported file type. Please upload a PDF, DOCX or TXT file.")
    return clean_text(text)


def clean_text(text: str) -> str:
    text = text.replace("\x00", "")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


# --------------------------------------------------------------------------
# Gemini call + response handling
# --------------------------------------------------------------------------
def build_prompt(resume_text: str, job_description: str) -> str:
    jd = job_description.strip()
    jd_block = (
        f"JOB DESCRIPTION:\n\"\"\"\n{jd[:8000]}\n\"\"\"\n"
        if jd
        else "JOB DESCRIPTION: (none provided - use general best practices)\n"
    )
    return (
        "Analyze this resume for ATS compatibility.\n\n"
        "Scoring: overall_score and every category score are integers from 0 to 100. "
        "Give 3-8 improvements ordered by priority (High first) and 2-4 rewrite_examples "
        "that take a real weak bullet from the resume and rewrite it with action verbs "
        "and measurable impact (do not invent facts; use placeholders like [X%] if needed). "
        "missing_keywords should be skills/terms an ATS would likely look for.\n\n"
        f"{jd_block}\n"
        f"RESUME:\n\"\"\"\n{resume_text[:MAX_RESUME_CHARS]}\n\"\"\""
    )


def parse_json_response(raw: str) -> dict:
    """Parse model output into a dict, tolerating markdown code fences."""
    if not raw or not raw.strip():
        raise ValueError("The model returned an empty response.")
    cleaned = raw.strip()
    cleaned = re.sub(r"^```(?:json)?\s*", "", cleaned)
    cleaned = re.sub(r"\s*```$", "", cleaned)
    try:
        return json.loads(cleaned)
    except json.JSONDecodeError:
        match = re.search(r"\{.*\}", cleaned, re.DOTALL)
        if match:
            return json.loads(match.group(0))
        raise ValueError("Could not parse the model's response as JSON.")


def _clamp(value, default=0) -> int:
    try:
        return max(0, min(100, int(round(float(value)))))
    except (TypeError, ValueError):
        return default


def normalize_result(data: dict) -> dict:
    """Make sure every expected field exists and has a sane type/range."""
    cats = data.get("category_scores") or {}
    result = {
        "overall_score": _clamp(data.get("overall_score")),
        "summary": str(data.get("summary", "")),
        "category_scores": {k: _clamp(cats.get(k)) for k in CATEGORY_LABELS},
        "strengths": [str(s) for s in (data.get("strengths") or [])],
        "missing_keywords": [str(k) for k in (data.get("missing_keywords") or [])],
        "improvements": [],
        "rewrite_examples": [],
    }
    order = {"high": 0, "medium": 1, "low": 2}
    for item in data.get("improvements") or []:
        if isinstance(item, dict):
            priority = str(item.get("priority", "Medium")).strip().title()
            if priority.lower() not in order:
                priority = "Medium"
            result["improvements"].append(
                {
                    "priority": priority,
                    "issue": str(item.get("issue", "")),
                    "suggestion": str(item.get("suggestion", "")),
                }
            )
    result["improvements"].sort(key=lambda i: order[i["priority"].lower()])
    for item in data.get("rewrite_examples") or []:
        if isinstance(item, dict) and item.get("original") and item.get("improved"):
            result["rewrite_examples"].append(
                {"original": str(item["original"]), "improved": str(item["improved"])}
            )
    return result


def analyze_resume(api_key: str, model: str, resume_text: str, job_description: str) -> dict:
    client = genai.Client(api_key=api_key)
    response = client.models.generate_content(
        model=model,
        contents=build_prompt(resume_text, job_description),
        config=types.GenerateContentConfig(
            system_instruction=SYSTEM_PROMPT,
            temperature=0.2,
            response_mime_type="application/json",
            response_schema=RESPONSE_SCHEMA,
        ),
    )
    return normalize_result(parse_json_response(response.text))


def score_label(score: int) -> str:
    if score >= 80:
        return "Excellent"
    if score >= 65:
        return "Good"
    if score >= 50:
        return "Needs work"
    return "Poor"


# --------------------------------------------------------------------------
# UI
# --------------------------------------------------------------------------
def get_api_key() -> str:
    try:
        key = st.secrets.get("GEMINI_API_KEY", "")
    except Exception:  # no secrets file locally
        key = ""
    return key or os.getenv("GEMINI_API_KEY", "")


def render_results(result: dict) -> None:
    score = result["overall_score"]
    st.divider()
    col1, col2 = st.columns([1, 3])
    with col1:
        st.metric("ATS Score", f"{score}/100", score_label(score), delta_color="off")
    with col2:
        st.progress(score / 100)
        st.write(result["summary"])

    st.subheader("Score breakdown")
    cols = st.columns(3)
    for i, (key, label) in enumerate(CATEGORY_LABELS.items()):
        with cols[i % 3]:
            st.metric(label, f"{result['category_scores'][key]}/100")

    left, right = st.columns(2)
    with left:
        st.subheader("Strengths")
        for s in result["strengths"] or ["No strengths listed."]:
            st.markdown(f"- {s}")
    with right:
        st.subheader("Missing keywords")
        if result["missing_keywords"]:
            st.write(", ".join(f"`{k}`" for k in result["missing_keywords"]))
        else:
            st.write("None identified.")

    st.subheader("Improvements")
    icons = {"High": "🔴", "Medium": "🟠", "Low": "🟢"}
    for item in result["improvements"]:
        with st.expander(f"{icons[item['priority']]} {item['priority']}: {item['issue']}"):
            st.write(item["suggestion"])

    if result["rewrite_examples"]:
        st.subheader("Example rewrites")
        for ex in result["rewrite_examples"]:
            st.markdown(f"**Before:** {ex['original']}")
            st.markdown(f"**After:** {ex['improved']}")
            st.write("")

    st.download_button(
        "Download report (JSON)",
        data=json.dumps(result, indent=2),
        file_name="ats_report.json",
        mime="application/json",
    )


def main() -> None:
    st.set_page_config(page_title="ATS Resume Checker", page_icon="📄", layout="wide")
    st.title("📄 ATS Resume Checker")
    st.caption("Upload your resume to get an ATS score and tips to improve it. Powered by Gemini Flash.")

    with st.sidebar:
        st.header("Settings")
        api_key = get_api_key()
        if api_key:
            st.success("API key loaded.")
        else:
            api_key = st.text_input("Gemini API key", type="password", help="Get one free at aistudio.google.com")
        model = st.text_input("Model", value=DEFAULT_MODEL)
        st.caption("Your resume is sent to Google's Gemini API for analysis.")

    uploaded = st.file_uploader("Resume (PDF, DOCX or TXT)", type=["pdf", "docx", "txt"])
    job_description = st.text_area(
        "Job description (optional, improves keyword matching)", height=160
    )

    if st.button("Analyze resume", type="primary"):
        if not api_key:
            st.error("Please provide a Gemini API key in the sidebar.")
            return
        if uploaded is None:
            st.error("Please upload a resume first.")
            return
        data = uploaded.getvalue()
        if len(data) > MAX_FILE_MB * 1024 * 1024:
            st.error(f"File is too large (max {MAX_FILE_MB} MB).")
            return

        try:
            text = extract_text(uploaded.name, data)
        except Exception as exc:
            st.error(f"Could not read the file: {exc}")
            return
        if len(text) < MIN_RESUME_CHARS:
            st.error(
                "Very little text could be extracted. If your resume is a scanned image, "
                "an ATS can't read it either - export a text-based PDF or DOCX instead."
            )
            return

        with st.spinner("Analyzing your resume..."):
            try:
                st.session_state["result"] = analyze_resume(
                    api_key, model.strip() or DEFAULT_MODEL, text, job_description
                )
            except Exception as exc:
                st.session_state.pop("result", None)
                st.error(f"Analysis failed: {exc}")
                return

    if "result" in st.session_state:
        render_results(st.session_state["result"])


if __name__ == "__main__":
    main()
